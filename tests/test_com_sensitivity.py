"""商业二维敏感性矩阵端点测试（P0-5，2026-10-05，D529）。

POST /api/com/sensitivity：退出 cap（±100bps，步长 50bps）×
租金增长率（±100bps，步长 50bps）→ 25 格 IRR / Equity Multiple，
全部复用 uw_commercial.compute_all()。
"""
from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from fastapi.testclient import TestClient

from app.main import app
from app.uw_commercial import compute_all, example_39_main

client = TestClient(app)


def _post(payload):
    r = client.post("/api/com/sensitivity", json=payload)
    assert r.status_code == 200, r.text[:300]
    return r.json()


def test_grid_shape_and_offsets():
    body = _post({"input": example_39_main()})
    assert body["exit_cap_offsets_bps"] == [-100, -50, 0, 50, 100]
    assert body["growth_offsets_bps"] == [-100, -50, 0, 50, 100]
    cells = body["cells"]
    assert len(cells) == 5 and all(len(row) == 5 for row in cells)
    for row in cells:
        for c in row:
            assert set(c.keys()) == {"irr", "equity_multiple"}
            # 终裁⑤：IRR 不收敛 → null（禁 0.0 冒充）
            assert c["irr"] is None or isinstance(c["irr"], (int, float))
            assert isinstance(c["equity_multiple"], (int, float))
    # 轴标签与基准一致
    assert len(body["exit_caps"]) == 5 and len(body["growths"]) == 5
    assert abs(body["exit_caps"][2] - body["exit_cap_base"]) < 1e-9
    assert abs(body["growths"][2] - body["growth_base"] / 100) < 1e-9
    assert abs(body["exit_caps"][4] - body["exit_caps"][0] - 0.02) < 1e-9


def test_center_cell_matches_baseline_compute_all():
    data = example_39_main()
    body = _post({"input": data})
    base = compute_all(data)
    base_exit = base["analysis"]["exit"]
    center = body["cells"][2][2]
    assert (center["irr"] is None and base_exit["irr"] is None) or \
        abs(center["irr"] - base_exit["irr"]) < 1e-3
    assert abs(center["equity_multiple"] - base_exit["equity_multiple"]) < 1e-3


def test_monotonicity_exit_cap_up_irr_down():
    # 退出 cap 越高 → 售价越低 → IRR 应单调不增（同行内比较）
    body = _post({"input": example_39_main()})
    for row in body["cells"]:
        irrs = [c["irr"] for c in row if c["irr"] is not None]
        for a, b in zip(irrs, irrs[1:]):
            assert a + 1e-9 >= b


def test_bad_input_400_not_500():
    r = client.post("/api/com/sensitivity", json={"input": "not-a-dict"})
    assert r.status_code == 400


def test_empty_input_ok():
    # 空输入不 500（compute_all 对空输入是良定义的）
    body = _post({"input": {}})
    assert len(body["cells"]) == 5


# ---------- Excel 敏感性矩阵 sheet（P0-5） ----------
import io as _io

from openpyxl import load_workbook as _load_wb

from app import excel_uw as _excel_uw


def _wb():
    data = example_39_main()
    r = compute_all(data)
    return _load_wb(_io.BytesIO(_excel_uw.build_uw_xlsx(r, data, [], {"has_rent_roll": True})))


def test_xlsx_sensitivity_sheet_exists():
    assert "敏感性矩阵" in _wb().sheetnames


def test_xlsx_sensitivity_cells_are_live_formulas():
    ws = _wb()["敏感性矩阵"]
    # 输入块：B2 基准退出 cap、B3 基准增长率、B4 持有年数（黄色可改格）
    assert ws["B2"].value == 0.05
    assert ws["B3"].value == 0.02
    assert ws["B4"].value == 5
    # IRR 网格 5×5 全部是 IRR 原生公式，且引用 $B$2 / $B$3 输入格
    irr_cells = [ws.cell(row=r, column=c).value for r in range(13, 18) for c in range(2, 7)]
    assert all(isinstance(v, str) and v.startswith("=IRR({") for v in irr_cells)
    # 公式引用轴头（增长率 $A<行> / 退出 cap <列>$12），轴头引用 $B$2 / $B$3 输入格
    assert all(f"$A{r}" in v and "$12" in v
               for r, v in zip([r for r in range(13, 18) for _ in range(5)], irr_cells))
    # EM 网格 5×5 全部是 SUMPRODUCT 公式
    em_cells = [ws.cell(row=r, column=c).value for r in range(20, 25) for c in range(2, 7)]
    assert all(isinstance(v, str) and v.startswith("=SUMPRODUCT(") for v in em_cells)
    # 轴头引用输入格（改黄色格即重算）
    assert str(ws["B12"].value).startswith("=$B$2")
    assert str(ws["A13"].value).startswith("=$B$3")
    # 基准格高亮（IRR 网格 D15、EM 网格 D22；B=首偏移列，D=±0 基准列）
    assert ws["D15"].fill.fgColor.rgb == "00D9EAD3"
    assert ws["D22"].fill.fgColor.rgb == "00D9EAD3"


def test_xlsx_sensitivity_formula_semantics_match_backend():
    # Excel 公式语义（纯 Python 镜像）必须等于后端中心格
    data = example_39_main()
    r = compute_all(data)
    ana, pro, ex = r["analysis"], r["proforma"], r["analysis"]["exit"]
    hold = int(ex["hold_years"]); noi0 = pro["noi"]; eq = ana["net_liquidity"]
    debt = ana["annual_debt_service"]
    abate = (data.get("analysis") or {}).get("abatements", 0)
    rem = ex["remaining_loan"]; fee = 0.04
    cap, g = ex["exit_cap_rate"], ex["noi_growth"]
    fl = [-eq]
    for y in range(1, hold + 1):
        noi = noi0 * (1 + g) ** (y - 1)
        cf = noi - abate - debt
        if y == hold:
            sale = noi0 * (1 + g) ** hold / cap
            cf += sale - rem - sale * fee
        fl.append(cf)
    from app.uw_commercial import _irr
    em = sum(f for f in fl[1:] if f > 0) / eq
    assert abs(_irr(fl) - ex["irr"]) < 1e-9
    assert abs(em - ex["equity_multiple"]) < 1e-9
