"""Phase 5 第四轮微修复回归测试（verify-then-fix）.

覆盖：
  A1: 否决态为独立第 4 状态（住宅"否决"/商业"VETO"，红）——"PASS · 否决态"零命中
  B2: PDF"关键假设"栏列出所有实际使用的假设（含默认值，标"默认"）；默认假设算出的数字旁加"*"
  B3: 退出净所得为"—"时，第 5 年现金流备注为"不含退出（退出假设缺失）"
  B4: 退出 cap 联动用户输入（市场 cap 显式填写时联动；否则 5% 标"默认"）；预测 CAP 率无数据时为"—"
  D8: Excel 总览 sheet 英文 label 中文化
  D9: verdict 理由中"score 总分"→"deal 评分"
"""
import io
import unittest

from openpyxl import load_workbook

from app import excel_uw, pdf_uw, uw_commercial, verdict


def _extract(pdf_bytes: bytes) -> str:
    import fitz
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    return "\n".join(p.get_text() for p in doc)


def _deal(**kw):
    base = {
        "name": "450 West 44th St", "address": "450 West 44th St, New York, NY 10036",
        "verdict": {"verdict": "PASS", "one_liner": "无 rent roll，DSCR 无法核保（待接数据）",
                    "reasons": ["一票否决：无 rent roll，DSCR 无法核保（待接数据）"]},
        "verdict_word": "VETO", "veto_count": 1,
        "dscr_trailing": None,
        "has_rent_roll": False,
        "assumptions": {"ask": 850000},
        "effective_assumptions": [
            {"key": "rate", "label": "年利率", "value": 6.5, "display": "6.5%", "is_default": True},
            {"key": "downPct", "label": "首付比例", "value": 30, "display": "30%", "is_default": True},
            {"key": "amort", "label": "摊销年数", "value": 25, "display": "25 年", "is_default": True},
            {"key": "vac", "label": "空置率", "value": 8, "display": "8%", "is_default": True},
            {"key": "capMkt", "label": "市场 cap", "value": 6, "display": "6%", "is_default": True},
            {"key": "hold_years", "label": "持有年数", "value": 5, "display": "5 年", "is_default": True},
            {"key": "exit_cap_rate", "label": "退出 cap", "value": 0.05,
             "display": "5.00%（默认）", "is_default": True},
        ],
        "comps": [],
    }
    base.update(kw)
    return base


class TestA1VetoFourthState(unittest.TestCase):
    def test_veto_word_is_veto_not_pass(self):
        d = uw_commercial.example_39_main()
        r = uw_commercial.compute_all(d)
        text = _extract(pdf_uw.build_uw_pdf(r, "enhanced", deal=_deal()))
        self.assertIn("VETO", text)
        self.assertIn("一票否决 ×1", text)
        # 否决态主词不许再是三档词
        self.assertNotIn("PASS · 否决态", text)
        self.assertNotIn("结论为否决态", text)

    def test_res_veto_word(self):
        # 住宅线主词为"否决"——verdict_word 由前端按同一规则生成，此处只验后端透传
        d = uw_commercial.example_39_main()
        r = uw_commercial.compute_all(d)
        deal = _deal(verdict_word="否决")
        text = _extract(pdf_uw.build_uw_pdf(r, "enhanced", deal=deal))
        self.assertIn("否决", text)
        self.assertNotIn("别碰 · 否决态", text)


class TestB2KeyAssumptions(unittest.TestCase):
    def test_defaults_listed_and_starred(self):
        d = uw_commercial.example_39_main()
        r = uw_commercial.compute_all(d)
        text = _extract(pdf_uw.build_uw_pdf(r, "enhanced", deal=_deal()))
        for label in ("年利率", "首付比例", "摊销年数", "空置率", "市场 cap", "持有年数", "退出 cap"):
            self.assertIn(label, text)
        # 默认值必须标注"默认"
        self.assertIn("6.5%（默认）", text)
        self.assertIn("5.00%（默认）", text)
        # 默认假设算出的数字旁加"*"
        self.assertIn("按默认假设测算", text)


class TestB3Year5Note(unittest.TestCase):
    def test_y5_note_when_no_exit(self):
        # 无 rent roll：tenants 为空 → projected_noi=0 → projected_resale=0 → resale_ok 假
        d = {"property": {"name": "T", "address": "T", "net_rentable_sf": 4500},
             "tenants": [], "vacant_sf": 0,
             "historical": {"vacancy_pct": 8}, "proforma": {"vacancy_pct": 8},
             "analysis": {"purchase_price": 850000, "down_pct": 30, "rate": 6.5,
                          "amort_years": 25}}
        r = uw_commercial.compute_all(d)
        text = _extract(pdf_uw.build_uw_pdf(r, "enhanced", deal=_deal()))
        # 退出净所得为"—"时，第 5 年不许写"含退出净所得"
        self.assertNotIn("含退出净所得", text)
        self.assertIn("不含退出（退出假设缺失）", text)

    def test_y5_note_when_exit_ok(self):
        # 有退出数据时保持"含退出净所得"
        d = uw_commercial.example_39_main()
        r = uw_commercial.compute_all(d)
        text = _extract(pdf_uw.build_uw_pdf(r, "enhanced", deal=_deal()))
        self.assertIn("含退出净所得", text)


class TestB4ExitCapLinkage(unittest.TestCase):
    def _exit(self, analysis):
        pro = {"noi": 200000, "total_potential": 300000, "total_expenses": 50000,
               "total_annual_lease": 0, "total_annual_uw": 0, "other_income": 0}
        return uw_commercial.compute_exit(analysis, pro, 595000, 0.065,
                                          "AMORTIZING", 25, 48210, 255000, 850000)

    def test_default_exit_cap(self):
        e = self._exit({"purchase_price": 850000})
        self.assertAlmostEqual(e["exit_cap_rate"], 0.05)
        self.assertEqual(e["exit_cap_source"], "默认")

    def test_market_cap_linkage(self):
        # 用户显式填了市场 cap 5.5% → 退出 cap 联动（pro-investor 投诉点）
        e = self._exit({"purchase_price": 850000, "market_cap_rate": 5.5,
                        "market_cap_rate_is_default": False})
        self.assertAlmostEqual(e["exit_cap_rate"], 0.055)
        self.assertEqual(e["exit_cap_source"], "联动市场cap")

    def test_frontend_default_market_cap_not_linked(self):
        # 前端默认值（用户没填）不触发联动，保持 5% 默认
        e = self._exit({"purchase_price": 850000, "market_cap_rate": 6,
                        "market_cap_rate_is_default": True})
        self.assertAlmostEqual(e["exit_cap_rate"], 0.05)
        self.assertEqual(e["exit_cap_source"], "默认")

    def test_explicit_exit_cap_wins(self):
        e = self._exit({"purchase_price": 850000, "exit_cap_rate": 7,
                        "market_cap_rate": 5.5, "market_cap_rate_is_default": False})
        self.assertAlmostEqual(e["exit_cap_rate"], 0.07)
        self.assertEqual(e["exit_cap_source"], "用户输入")

    def test_predicted_cap_rate_dash(self):
        d = uw_commercial.example_39_main()
        r = uw_commercial.compute_all(d)
        text = _extract(pdf_uw.build_uw_pdf(r, "enhanced", deal=_deal()))
        self.assertNotIn("0.00%", text)


class TestD8ExcelLabels(unittest.TestCase):
    def _labels(self, ws, n=80):
        return [ws.cell(row=i, column=1).value for i in range(1, n)]

    def test_com_overview_no_english_labels(self):
        r = uw_commercial.example_39_main()
        res = uw_commercial.compute_all(r)
        x = excel_uw.build_uw_xlsx(res, {}, comps=[], meta={"has_rent_roll": True})
        ws = load_workbook(io.BytesIO(x))["总览"]
        labels = " | ".join(str(v) for v in self._labels(ws) if v)
        for eng in ("verdict", "Confidence"):
            self.assertNotIn(eng, labels)
        self.assertIn("结论", labels)
        self.assertIn("偿债覆盖率（DSCR，历史主）", labels)

    def test_res_overview_no_english_labels(self):
        score = {"total": 43, "metrics": {"cash_flow_monthly": -19000, "cash_on_cash": -0.05,
                                          "piti": 7000, "dscr": 0.5, "cash_to_close": 870000,
                                          "egi": 8000, "opex": 1000}}
        x = excel_uw.build_res_xlsx(score, {"price": 3000000, "vacancy_pct": 8},
                                    "450 West 44th St",
                                    verdict={"verdict": "别碰", "one_liner": "测试"},
                                    confidence={"score": 65}, comps=[])
        ws = load_workbook(io.BytesIO(x))["总览"]
        labels = " | ".join(str(v) for v in self._labels(ws) if v)
        for eng in ("verdict", "Confidence", "CoC", "PITI"):
            self.assertNotIn(eng, labels)
        self.assertIn("结论", labels)
        self.assertIn("置信度", labels)
        self.assertIn("偿债覆盖率（DSCR）", labels)


class TestD9DealScoreNaming(unittest.TestCase):
    def test_score_renamed(self):
        v = verdict.verdict("residential", {"total": 43, "metrics": {}}, 65, {})
        reasons = " | ".join(v.get("reasons") or [])
        self.assertIn("deal 评分", reasons)
        self.assertNotIn("score 总分", reasons)


if __name__ == "__main__":
    unittest.main()
