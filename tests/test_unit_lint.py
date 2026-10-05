"""商业核保数据标准 v1.0 §5 —— LINT-01…08 单元级门禁（pytest）.

这些是"口径守卫"：防止 P0-1 类单位 bug reintroduce。任一 FAIL = 口径事故。
运行前按铁律：rm -f dealdesk.db（本目录无 DB 依赖，但保持一致）。
"""

import ast
import io
import re
from pathlib import Path

import pytest

from app import uw_commercial
from app.uw_commercial import (
    DEFAULT_EXIT_CAP_PCT,
    DEFAULT_NOI_GROWTH_PCT,
    EXIT_FEE_PCT,
    TIER_A,
    TIER_S,
    compute_all,
    default_inputs,
    normalize_pct,
)

APP = Path(__file__).resolve().parent.parent / "app"
STATIC = Path(__file__).resolve().parent.parent / "static"


# ---------------------------------------------------------------- LINT-01
class TestLint01PctMultiply:
    """百分制字段进乘法/除法无 normalize 判 FAIL.

    规则（ast 扫描 app/uw_commercial.py）：
    (a) 原始输入形态 —— .get("vacancy_pct") / ["down_pct"] / .market_cap_rate 等
        直接参与 * / 运算，必须被 normalize_pct(...) 包裹；
    (b) 以 _pct 结尾的局部变量名或裸名 rate / noi_growth / exit_cap 参与 * / 运算，
        必须在同一函数内经 normalize_pct(...) 赋值过（归一化点唯一）；
    (c) 全大写命名常量（如 EXIT_FEE_PCT / 100.0）豁免 —— 唯一允许的手工 /100。
    """

    _SUSPECT_KEYS = ("vacancy_pct", "down_pct", "market_cap_rate", "exit_cap_rate",
                     "noi_growth")
    _SUSPECT_NAMES = _SUSPECT_KEYS + ("rate", "noi_growth", "exit_cap")

    def _raw_pct_operand(self, node):
        """操作数是否为原始百分制输入形态？"""
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) \
                and node.func.attr == "get" and node.args \
                and isinstance(node.args[0], ast.Constant) \
                and node.args[0].value in self._SUSPECT_KEYS:
            return f'.get("{node.args[0].value}")'
        if isinstance(node, ast.Subscript) and isinstance(node.slice, ast.Constant) \
                and node.slice.value in self._SUSPECT_KEYS:
            return f'["{node.slice.value}"]'
        if isinstance(node, ast.Attribute) and node.attr in self._SUSPECT_KEYS:
            return f".{node.attr}"
        return None

    def test_no_raw_pct_in_arithmetic(self):
        src = (APP / "uw_commercial.py").read_text()
        tree = ast.parse(src)
        for node in ast.walk(tree):
            for child in ast.iter_child_nodes(node):
                child._parent = node
        offenses = []
        for fn in [n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef)]:
            # 本函数内经 normalize_pct 赋值的局部名 = 已归一化
            normalized = set()
            for sub in ast.walk(fn):
                if isinstance(sub, ast.Assign) and isinstance(sub.value, ast.Call) \
                        and getattr(sub.value.func, "id", "") == "normalize_pct":
                    for t in sub.targets:
                        if isinstance(t, ast.Name):
                            normalized.add(t.id)
            for sub in ast.walk(fn):
                if not isinstance(sub, ast.BinOp) or not isinstance(sub.op, (ast.Mult, ast.Div)):
                    continue
                # normalize_pct 定义体内部的 /100 是唯一的合法手工换算
                if fn.name == "normalize_pct":
                    continue
                # (c) 全大写常量豁免
                names = {n.id for n in ast.walk(sub) if isinstance(n, ast.Name)}
                if names and all(n.isupper() for n in names):
                    continue
                for operand in (sub.left, sub.right):
                    for cand in ast.walk(operand):
                        raw = self._raw_pct_operand(cand)
                        if raw:
                            # 必须被 normalize_pct(...) 调用包裹
                            p, wrapped = getattr(cand, "_parent", None), False
                            while p is not None and p is not fn:
                                if isinstance(p, ast.Call) and getattr(p.func, "id", "") == "normalize_pct":
                                    wrapped = True
                                    break
                                p = getattr(p, "_parent", None)
                            if not wrapped:
                                offenses.append(
                                    f"{fn.name}:{sub.lineno} 原始输入 {raw} 未归一化进乘除")
                        elif isinstance(cand, ast.Name) and cand.id in self._SUSPECT_NAMES \
                                and cand.id not in normalized and not cand.id.isupper():
                            offenses.append(
                                f"{fn.name}:{sub.lineno} 变量 {cand.id} 未归一化进乘除")
        # 去重
        offenses = sorted(set(offenses))
        assert not offenses, "LINT-01 FAIL:\n" + "\n".join(offenses)

    def test_normalize_pct_exists_and_shape(self):
        assert normalize_pct(8) == pytest.approx(0.08)
        assert normalize_pct("6.5") == pytest.approx(0.065)
        assert normalize_pct(None) == 0.0
        assert normalize_pct("abc") == 0.0


# ---------------------------------------------------------------- LINT-02
class TestLint02Registry:
    """注册表 / 默认值扫描：所有 *_pct 字段必须在 TIER_S/TIER_A 注册表中；
    模板默认值必须为百分制（>1 才是百分制数字，0 允许）。"""

    def test_pct_fields_registered(self):
        pct_fields = {"vacancy_pct", "down_pct", "rate", "market_cap_rate",
                      "exit_cap_rate", "noi_growth"}
        registered = set(TIER_S) | set(TIER_A)
        assert pct_fields <= registered, f"未注册：{pct_fields - registered}"

    def test_template_defaults_are_pct(self):
        d = default_inputs()
        a = d["analysis"]
        for k, v in (("down_pct", a["down_pct"]), ("rate", a["rate"]),
                     ("market_cap_rate", a["market_cap_rate"]),
                     ("noi_growth", a["noi_growth"]),
                     ("exit_cap_rate", a["exit_cap_rate"])):
            assert v == 0 or v > 1, f"LINT-02 FAIL：模板默认 {k}={v} 疑似小数（百分制应 >1 或 0）"
        assert a["down_pct"] == 30 and a["rate"] == 6.5

    def test_named_constants_no_magic(self):
        assert EXIT_FEE_PCT == 4.0
        assert DEFAULT_NOI_GROWTH_PCT == 2.0
        assert DEFAULT_EXIT_CAP_PCT == 5.0


# ---------------------------------------------------------------- LINT-03
class TestLint03XlsxPctFormat:
    """Excel PCT 格式格的底层值必须是小数（0-1）；百分制数字进 PCT 格式格判 FAIL."""

    def test_pct_cells_hold_decimals(self):
        from openpyxl import load_workbook
        from app import excel_uw
        d = uw_commercial.example_39_main()
        r = compute_all(d)
        x = excel_uw.build_uw_xlsx(r, d, comps=[], meta={"has_rent_roll": True})
        wb = load_workbook(io.BytesIO(x))
        bad = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.value is None or isinstance(c.value, str):
                        continue
                    fmt = (c.number_format or "").lower()
                    if "0.00%" in fmt or fmt.endswith("%"):
                        v = c.value
                        if isinstance(v, (int, float)) and not (0 <= abs(v) <= 1.5):
                            bad.append(f"{ws.title}!{c.coordinate}={v} 格式={c.number_format}")
        assert not bad, "LINT-03 FAIL：PCT 格式格出现百分制裸值：\n" + "\n".join(bad[:10])


# ---------------------------------------------------------------- LINT-04
class TestLint04HonestFormatters:
    """formatter 诚实性：NaN/None/缺失必须渲染 N/A 标记（—），不许 0；
    源码中禁止 `|| 0` 回退缺失的 noi."""

    def test_pdf_formatters_dash_not_zero(self):
        from app import pdf_uw
        for fn in (pdf_uw._money, pdf_uw._pct, pdf_uw._num):
            assert fn(None) == "—"
            assert fn(float("nan")) == "—"
        assert pdf_uw._money(0) != "—"  # 真零可渲染为零

    def test_no_noi_or_zero_fallback_in_sources(self):
        bad = []
        for p in (APP / "workbench.py", STATIC / "d.js", APP / "main.py"):
            src = p.read_text()
            for i, line in enumerate(src.splitlines(), 1):
                if re.search(r"noi_annual\s*\|\|\s*0", line):
                    bad.append(f"{p.name}:{i}: {line.strip()}")
        assert not bad, "LINT-04 FAIL：发现 noi || 0 回退：\n" + "\n".join(bad)

    def test_irr_nonconvergent_is_none_not_zero(self):
        assert uw_commercial._irr([100, 1, 1, 1]) is None  # 无符号变化 → 不收敛
        e = uw_commercial.compute_exit(
            {"purchase_price": 0},
            {"noi": 0, "total_potential": 1, "total_expenses": 0},
            0, 0.06, "IO", 30, 0, 0, 0)
        assert e["irr"] is None and e["sale_price"] is None, \
            "LINT-04 FAIL：退化输入的 IRR/出售价不许用 0.0 冒充"


# ---------------------------------------------------------------- LINT-05
class TestLint05ScenarioIsolation:
    """双情景变量隔离：historical 改 vacancy 不许污染 proforma（反之亦然）."""

    def test_scenarios_isolated(self):
        d = default_inputs()
        d["historical"] = {"vacancy_pct": 8, "other_income": 1000000}
        d["proforma"] = {"vacancy_pct": 2, "other_income": 1000000}
        r = compute_all(d)
        assert r["historical"]["vacancy_loss"] == pytest.approx(80000)
        assert r["proforma"]["vacancy_loss"] == pytest.approx(20000)
        # 输入字典本身不许被 compute 过程改写
        assert d["historical"]["vacancy_pct"] == 8
        assert d["proforma"]["vacancy_pct"] == 2


# ---------------------------------------------------------------- LINT-06
class TestLint06NoMagicNumbers:
    """魔法数字扫描：uw_commercial.py 内裸 0.04/0.02/0.05（退出费/增长率/退出cap）
    必须经命名常量。"""

    def test_no_bare_magic(self):
        src = (APP / "uw_commercial.py").read_text()
        tree = ast.parse(src)
        const_names = {"EXIT_FEE_PCT", "DEFAULT_NOI_GROWTH_PCT", "DEFAULT_EXIT_CAP_PCT",
                       "DEFAULT_HOLD_YEARS"}
        assigned = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Assign):
                for t in node.targets:
                    if isinstance(t, ast.Name) and t.id in const_names:
                        assigned.add(t.id)
        assert assigned == const_names, f"命名常量缺失：{const_names - assigned}"
        # 业务逻辑行不许出现裸魔法数字（注释/常量定义行除外）
        bad = []
        for i, line in enumerate(src.splitlines(), 1):
            code = line.split("#")[0]
            if re.search(r"[^0-9a-zA-Z_](0\.04|0\.02|0\.05)([^0-9]|$)", code):
                # 允许：常量定义行、_div 等工具函数、测试口径注释行
                if any(k in line for k in const_names) or "def " in line:
                    continue
                bad.append(f"uw_commercial.py:{i}: {line.strip()}")
        assert not bad, "LINT-06 FAIL：裸魔法数字：\n" + "\n".join(bad)


# ---------------------------------------------------------------- LINT-07
class TestLint07PmtConvention:
    """PMT 月供口径：amort_payment 必须是月复利等额本息（PMT），名义年利率按月计息."""

    def test_monthly_amortizing(self):
        from app.uw_commercial import amort_payment
        pmt = amort_payment(1_000_000, 0.06, 30)
        # 年债务偿还额 = 月复利等额本息月供 × 12：100万/6%/30年 ≈ 71,946/年
        assert pmt == pytest.approx(71946, rel=0.002), pmt
        # 旧口径（按年复利一次付）≈ 72,649 —— 不许回归
        assert abs(pmt - 1_000_000 * 0.06 / (1 - 1.06 ** -30)) > 500

    def test_compute_matches_pmt(self):
        d = default_inputs()
        d["analysis"] = {"purchase_price": 1000000, "down_pct": 30, "rate": 6,
                         "amort_years": 30, "amort_type": "AMORTIZING"}
        r = compute_all(d)
        assert r["analysis"]["annual_debt_service"] == pytest.approx(
            700000 * 0.06 / 12 / (1 - (1 + 0.06 / 12) ** -360) * 12, rel=1e-6)


# ---------------------------------------------------------------- LINT-08
class TestLint08TierRegistry:
    """Tier S/A 注册表：compute_all 输出必须带 data_gaps；Tier S 缺失 → tier_s_ok False."""

    def test_data_gaps_present(self):
        r = compute_all(default_inputs())
        assert "data_gaps" in r and "tier_s" in r["data_gaps"] and "tier_a" in r["data_gaps"]
        assert r["tier_s_ok"] is False  # blank 模板缺 purchase_price/sf
        assert "purchase_price" in r["data_gaps"]["tier_s"]

    def test_tier_s_ok_when_complete(self):
        d = uw_commercial.example_39_main()
        r = compute_all(d)
        assert r["tier_s_ok"] is True, r["data_gaps"]

    def test_tier_a_gap_tracked_not_silent(self):
        d = uw_commercial.example_39_main()
        del d["analysis"]["market_cap_rate"]
        r = compute_all(d)
        assert "analysis.market_cap_rate" in r["data_gaps"]["tier_a"]
