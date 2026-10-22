"""机构级导出：带公式的承销模型 Excel 与机构投委会备忘录 PDF。

核心保证：Excel 公式在不改输入时与引擎结果逐项一致（装了 LibreOffice 时真重算核对），改关键假设会在 Excel 里重算；
备忘录各章节齐全、数字与引擎同源、缺数不编；敏感性每格真重跑引擎，基准格等于基准结果。
"""
from __future__ import annotations

import copy
import io
import itertools
import shutil
import subprocess

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.domain.v2 import kill_criteria
from app.domain.v2.cre import engine
from app.domain.v2.cre.example import suburban_office
from app.domain.v2.model import run_model
from app.services.render import cre_excel, cre_memo

META = {"name": "Suburban Office", "address": "1 Park Ave", "version_no": 2, "date": "2026-10-22"}


@pytest.fixture(scope="module")
def case():
    cre = engine.parse(suburban_office()).model_dump()
    return cre, engine.run(cre)


def _recalc(xlsx: bytes, tmp_path) -> object:
    src = tmp_path / "m.xlsx"
    src.write_bytes(xlsx)
    out = tmp_path / "out"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=180)
    return load_workbook(out / "m.xlsx", data_only=True)


def test_workbook_has_formulas_wired_to_inputs(case):
    cre, r = case
    wb = load_workbook(io.BytesIO(cre_excel.build_cre_xlsx(cre, r, META)))
    assert wb.sheetnames == ["假设", "现金流", "收益", "租金表", "月度", "瀑布", "说明"]
    a, cf, ret = wb["假设"], wb["现金流"], wb["收益"]
    assert a["B4"].value == 24_000_000 and a["B6"].value == pytest.approx(0.0775)
    noi = cre_excel.ROW["noi"]
    assert cf[f"B{noi}"].value == f"=B{cre_excel.ROW['egr']}+B{cre_excel.ROW['opex']}"
    assert ret["B5"].value == "=B3/B4"                     # 售价 = 退出 NOI ÷ 退出资本化率
    assert ret["B4"].value == "='假设'!$B$6"
    assert ret["B16"].value.startswith("=IRR(B12:")
    assert ret["B12"].value == "=-'假设'!$B$4*(1+'假设'!$B$5)"
    # 第 N+1 年只用于退出 NOI：不计还本付息与 DSCR
    last = cre_excel._col(cre["hold_years"] + 2)
    assert cf[f"{last}{cre_excel.ROW['debt_service']}"].value is None
    # 假设续上的租约段在租金表里如实标注
    kinds = {row[6] for row in wb["租金表"].iter_rows(min_row=2, values_only=True)}
    assert kinds == {"已签约", "假设续上（MLA）"}


def test_engine_lines_add_up_to_formula_rows(case):
    """Excel 的 EGR / 费用合计 / NOI / 债前现金流公式就是引擎科目相加：用引擎年度值验恒等式（不依赖 Excel）。"""
    _, r = case
    a = r["annual"]
    for y in range(len(a["noi"])):
        egr = a["base_rent"][y] + a["free_rent"][y] + a["recoveries"][y] + a["general_vacancy"][y] + a["credit_loss"][y]
        opex = a["opex_recoverable"][y] + a["opex_non_recoverable"][y] + a["mgmt_fee"][y]
        assert egr == pytest.approx(a["egr"][y], abs=1e-6) and opex == pytest.approx(a["opex"][y], abs=1e-6)
        assert egr + opex == pytest.approx(a["noi"][y], abs=1e-6)
        cfbd = a["noi"][y] + a["ti"][y] + a["lc"][y] + a["reserves"][y]
        assert cfbd == pytest.approx(a["cf_before_debt"][y], abs=1e-6)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="没装 LibreOffice，无法重算公式")
def test_libreoffice_recalc_matches_engine_and_reacts_to_inputs(case, tmp_path):
    cre, r = case
    wb = _recalc(cre_excel.build_cre_xlsx(cre, r, META), tmp_path)
    ret = wb["收益"]
    for row in [*range(3, 10), *range(16, 31)]:  # 退出块与指标块（B = 公式，C = 引擎）
        label, formula_v, engine_v = (ret.cell(row=row, column=c).value for c in (1, 2, 3))
        if isinstance(formula_v, (int, float)) and isinstance(engine_v, (int, float)):
            assert formula_v == pytest.approx(engine_v, rel=1e-6, abs=1e-6), label
    assert ret["B17"].value == pytest.approx(r["metrics"]["levered_irr_pct"] / 100, abs=1e-8)
    wf = wb["瀑布"]
    checks = [wf.cell(row=i, column=5).value for i in range(5, 5 + cre["hold_years"] + 1)]
    assert all(abs(x) < 1e-6 for x in checks)
    # 改退出资本化率：售价与 IRR 在 Excel 里重算，方向正确
    src = load_workbook(io.BytesIO(cre_excel.build_cre_xlsx(cre, r, META)))
    src["假设"]["B6"].value = 0.0875
    buf = io.BytesIO()
    src.save(buf)
    (tmp_path / "b").mkdir()
    wb2 = _recalc(buf.getvalue(), tmp_path / "b")
    assert wb2["收益"]["B5"].value == pytest.approx(r["metrics"]["forward_noi"] / 0.0875, rel=1e-9)
    assert wb2["收益"]["B17"].value < ret["B17"].value


def _pdf_text(b: bytes) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(b)).pages)


def test_memo_sections_and_numbers(case):
    cre, r = case
    kill = kill_criteria.evaluate(kill_criteria.DEFAULT_INSTITUTIONAL, run_model("institutional", suburban_office())["metrics"])
    b = cre_memo.build_cre_memo(cre, r, META, kill)
    assert b[:5] == b"%PDF-"
    text = _pdf_text(b)
    for heading in ("交易概要", "租金表", "逐年现金流", "资本结构", "GP / LP", "敏感性", "主要假设", "方法与局限"):
        assert heading in text, heading
    assert "16.90%" in text and "11.85%" in text and "$1,855,703" in text   # 有杠杆 / 无杠杆 IRR、promote
    assert "Northwind Legal" in text and "通过" in text


def test_memo_without_debt_or_waterfall_does_not_invent_numbers():
    inp = {**suburban_office(), "debt": None, "waterfall": None}
    cre = engine.parse(inp).model_dump()
    r = engine.run(cre)
    text = _pdf_text(cre_memo.build_cre_memo(cre, r, META, None))
    assert "无（全股权）" in text and "未设置分配瀑布" in text and "没有设置淘汰条件" in text
    wb = load_workbook(io.BytesIO(cre_excel.build_cre_xlsx(cre, r, META)))
    assert "瀑布" not in wb.sheetnames


def test_sensitivity_reruns_engine_and_centres_on_base(case):
    cre, r = case
    s = cre_memo.sensitivity(cre)
    g = s["cap_rent"]
    i, j = g["rows"].index(cre["exit_cap_pct"]), g["cols"].index(0.0)
    assert g["grid"][i][j] == pytest.approx(r["metrics"]["levered_irr_pct"], abs=1e-9)
    assert all(a < b for a, b in itertools.pairwise(g["grid"][i]))                             # 租金越高 IRR 越高
    assert g["grid"][0][j] > g["grid"][-1][j]                                                 # 退出 cap 越高 IRR 越低
    rd = s["renew_down"]
    assert cre["market"]["renewal_prob_pct"] in rd["rows"] and cre["market"]["downtime_months"] in rd["cols"]
    k = rd["rows"].index(cre["market"]["renewal_prob_pct"])
    assert rd["grid"][k][rd["cols"].index(cre["market"]["downtime_months"])] == pytest.approx(
        r["metrics"]["levered_irr_pct"], abs=1e-9)


def test_walt_expiry_and_exit_year_drag(case):
    cre, r = case
    # 剩余月数：Northwind 2031-03 → 63 个月，Contoso 2027-06 → 18，Fabrikam 2035-06 → 114（分析起点 2026-01）
    exp = (32_000 * 63 + 28_000 * 18 + 22_000 * 114) / 82_000 / 12
    assert cre_memo.walt_years(cre) == pytest.approx(exp)
    sched = dict(cre_memo.expiry_schedule(cre))
    assert sched[2] == 28_000 and sched[6] == 32_000 and sum(sched.values()) == 60_000
    c = copy.deepcopy(cre)
    c["market"].update(renewal_prob_pct=50, downtime_months=3)
    r2 = engine.run(c)
    drag = cre_memo.exit_year_drag(c, r2)
    a = r2["annual"]
    assert drag > 0 and drag == pytest.approx(-(a["turnover_vacancy"][7] + a["free_rent"][7]))
    assert cre_memo.exit_year_drag(cre, r) == pytest.approx(-(r["annual"]["turnover_vacancy"][7]
                                                              + r["annual"]["free_rent"][7]))


def test_missing_values_render_as_dash():
    assert cre_memo.pct(None) == "—" and cre_memo.money(None) == "—" and cre_memo.mult(None) == "—"
    assert cre_memo.money(-1234.4) == "($1,234)"


def test_stabilized_exit_basis_flows_into_both_exports():
    cre = engine.parse({**suburban_office(), "exit_noi_basis": "stabilized"}).model_dump()
    r = engine.run(cre)
    wb = load_workbook(io.BytesIO(cre_excel.build_cre_xlsx(cre, r, META)))
    ret = wb["收益"]
    assert ret["A3"].value.startswith("退出年稳定化 NOI")
    assert ret["B3"].value == f"={r['metrics']['forward_noi_stabilized']}"   # 稳定化 NOI 是引擎值，售价公式照旧引用它
    assert ret["B5"].value == "=B3/B4"
    text = _pdf_text(cre_memo.build_cre_memo(cre, r, META, None))
    assert "稳定化" in text
