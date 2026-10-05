"""Phase 5 第二轮修复回归测试（panel2 6 人复测残留问题）.

覆盖清单 A–E：
  item 1: verdict 否决态显示词（后端 veto 消息同源；前端显示词由 d.js verdictDisplay 生成，此处测后端消息）
  item 3: 无 rent roll 时 low_dscr 否决理由 = "无 rent roll，DSCR 无法核保（待接数据）"，不许 0.00
  item 4: 中位 $/SF 真中位数；valuate_comps 偶数个时 median_adjusted 为真中位数
  item 5: 年还本付息按月摊还（银行标准）；$2.24M/6.5%/25年 = $181,495.68（broker 实证）
  item 6: Excel DSCR 分子用银行口径 NOI；B13 为 PMT 原生公式（按月摊还）
  item 7: PDF 增强版——否决态徽标同行、USPAP 声明、总费用$0 标注、无意义 ROI 行删除、假设中文标签
  item 12: PDF/Excel 假设区中文标签
  item 14: _parse_rr_date 支持 "Jan 1, 2025"
"""
import os
import sys
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app import excel_uw, pdf_uw, scoring_commercial, uw_commercial, workbench  # noqa: E402
from app.finance import monthly_payment  # noqa: E402
from app.main import _parse_rr_date  # noqa: E402
from app.workbench import WbComp, valuate_comps  # noqa: E402


def approx(a, b, tol=0.01):
    return abs(a - b) <= tol


class TestMonthlyAmortization(unittest.TestCase):
    def test_amort_payment_is_monthly_based(self):
        # broker 实证：$2.24M/6.5%/25年 按年 $183,638.52 vs 按月 $181,495.68
        pmt = uw_commercial.amort_payment(2240000, 0.065, 25)
        self.assertTrue(approx(pmt, 181495.68, 1.0), f"got {pmt}")
        self.assertTrue(approx(pmt, monthly_payment(2240000, 6.5, 25) * 12, 0.01))
        # 不再是旧的按年复利值
        self.assertFalse(approx(pmt, 183638.52, 1.0), "仍是按年摊还口径！")

    def test_remaining_balance_monthly_consistent(self):
        from app.finance import remaining_balance as _rbm
        pmt = uw_commercial.amort_payment(29400000, 0.06, 30)
        rem = uw_commercial.remaining_balance(29400000, 0.06, 30, pmt, 5)
        self.assertTrue(0 < rem < 29400000)
        # 与 finance.monthly 口径一致（payment 参数兼容保留，不再使用）
        self.assertTrue(approx(rem, _rbm(29400000, 6.0, 30, 5), 0.01))


class TestTrueMedian(unittest.TestCase):
    def test_valuate_comps_even_count_true_median(self):
        # panel2 remote ground truth：10 个 comp，$/SF 排序后第 6 个=771，真中位=739.5
        psfs = [147, 300, 450, 600, 708, 771, 900, 1050, 1200, 1390]
        comps = [WbComp(address=f"c{i}", price=p * 4500, sf=4500,
                         sale_date="2026-01-15", distance_miles=0.5,
                         adjustment_pct=0, source="TopHap CMA")
                 for i, p in enumerate(psfs)]
        d = valuate_comps(comps)
        # 调整价中位数（真中位数，非上中位数）
        adj = sorted(p * 4500 for p in psfs)
        expect = (adj[4] + adj[5]) / 2
        self.assertTrue(approx(d["median_adjusted"], expect, 0.01),
                        f"got {d['median_adjusted']}, expect {expect}")
        self.assertFalse(approx(d["median_adjusted"], adj[5], 0.01),
                         "仍是上中位数（旧 bug）！")


class TestVetoMessageNoRentRoll(unittest.TestCase):
    def _base(self):
        return {
            "asset_class": "retail_strip", "price": 4350000,
            "annual_base_rent": 0, "other_income_annual": 0,
            "vacancy_pct": 8, "tax_annual": 0, "insurance_annual": 0,
            "mgmt_pct": 8, "maint_pct": 5, "property_age_years": 0,
            "utilities_cam_gap_annual": 0, "building_sf": 4500, "units": 0,
            "tranches": [{"balance": 3045000, "rate": 6.5, "rate_type": "fixed",
                          "term_years": 25, "balloon_years": None}],
            "structure": "standard", "phase1_clear": True, "zoning_ok": True,
        }

    def test_no_rent_roll_veto_message(self):
        d = self._base()
        d["has_rent_roll"] = False
        m = scoring_commercial.compute_metrics(d)
        vetoes = scoring_commercial.check_vetoes(d, m)
        low = [v for v in vetoes if v["code"] == "low_dscr"]
        self.assertTrue(low, "应触发 low_dscr 否决")
        self.assertEqual(low[0]["message"], "无 rent roll，DSCR 无法核保（待接数据）")
        self.assertNotIn("0.00", low[0]["message"], "不许出现 0.00x 当硬数字")

    def test_with_rent_roll_veto_message_numeric(self):
        d = self._base()
        d["has_rent_roll"] = True
        d["annual_base_rent"] = 240000
        m = scoring_commercial.compute_metrics(d)
        vetoes = scoring_commercial.check_vetoes(d, m)
        low = [v for v in vetoes if v["code"] == "low_dscr"]
        self.assertTrue(low)
        self.assertIn("DSCR", low[0]["message"])
        self.assertNotIn("待接数据", low[0]["message"])


class TestRrDateParse(unittest.TestCase):
    def test_jan_text_date(self):
        self.assertEqual(_parse_rr_date("Jan 1, 2025"), "2025-01-01")
        self.assertEqual(_parse_rr_date("January 1, 2025"), "2025-01-01")
        self.assertEqual(_parse_rr_date("Dec 31, 2030"), "2030-12-31")

    def test_existing_formats_still_work(self):
        self.assertEqual(_parse_rr_date("2023-06-15"), "2023-06-15")
        self.assertEqual(_parse_rr_date("2023/06/15"), "2023-06-15")
        self.assertEqual(_parse_rr_date("not a date"), "")


class TestExcelBankNoi(unittest.TestCase):
    def _result(self):
        d = uw_commercial.example_39_main()
        return uw_commercial.compute_all(d), d

    def test_dscr_uses_bank_noi_and_pmt_formula(self):
        from openpyxl import load_workbook
        import io as _io
        r, data = self._result()
        xbytes = excel_uw.build_uw_xlsx(r, data, comps=[], meta={"bank_noi": 169000.0, "has_rent_roll": True})
        wb = load_workbook(_io.BytesIO(xbytes))
        ws = wb["总览"]
        labels = {ws.cell(row=i, column=1).value: i for i in range(1, 30)}
        # 银行口径 NOI 行存在
        self.assertIn("NOI（银行口径）$", labels)
        noi_row = labels["NOI（银行口径）$"]
        self.assertEqual(ws.cell(row=noi_row, column=2).value, 169000.0)
        # DSCR 公式引用银行口径 NOI 行
        dscr_row = labels["偿债覆盖率（DSCR，历史主）"]
        formula = ws.cell(row=dscr_row, column=2).value
        self.assertIn(f"B{noi_row}", str(formula))
        # B13（年还本付息）是 PMT 原生公式，不是硬编码
        ds_row = labels["年还本付息 $"]
        ds_formula = ws.cell(row=ds_row, column=2).value
        self.assertTrue(str(ds_formula).startswith("=PMT("), f"got {ds_formula}")
        self.assertIn("/12", str(ds_formula), "PMT 必须按月摊还")

    def test_no_bank_noi_fallback(self):
        from openpyxl import load_workbook
        import io as _io
        r, data = self._result()
        xbytes = excel_uw.build_uw_xlsx(r, data, comps=[], meta={})
        wb = load_workbook(_io.BytesIO(xbytes))
        ws = wb["总览"]
        labels = [ws.cell(row=i, column=1).value for i in range(1, 30)]
        self.assertIn("NOI（历史/银行口径前）$", labels)

    def test_assumption_sheet_chinese_labels(self):
        from openpyxl import load_workbook
        import io as _io
        r, data = self._result()
        xbytes = excel_uw.build_uw_xlsx(r, data, comps=[], meta={})
        wb = load_workbook(_io.BytesIO(xbytes))
        ax = wb["假设"]
        labels = [ax.cell(row=i, column=1).value for i in range(2, 12)]
        labels = [x for x in labels if x]
        self.assertTrue(labels, "假设表为空")
        for lab in labels:
            self.assertNotRegex(lab, r"^[a-z_]+$",
                                f"英文内部字段名泄露：{lab}")


class TestEnhancedPdfVeto(unittest.TestCase):
    def _extract(self, pdf_bytes: bytes) -> str:
        import fitz
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        return "\n".join(p.get_text() for p in doc)

    def _deal(self, **kw):
        base = {
            "name": "450 West 44th St", "address": "450 West 44th St, New York, NY 10036",
            "verdict": {"verdict": "PASS", "one_liner": "DSCR 0.93 < 1.0（生死线），贴钱持有",
                        "reasons": ["一票否决：DSCR 0.93 < 1.25，直接否决不参与打分"]},
            "verdict_word": "VETO", "veto_count": 1,  # Phase 5 第四轮 A1：否决态独立第 4 状态
            "dscr_trailing": 0.93,
            "assumptions": {"structure": "standard", "ask": 4350000,
                            "exitStrategy": "", "rate": 6.5},
            "comps": [],
        }
        base.update(kw)
        return base

    def test_veto_word_and_badge_same_line(self):
        d = uw_commercial.example_39_main()
        r = uw_commercial.compute_all(d)
        pdf = pdf_uw.build_uw_pdf(r, "enhanced", deal=self._deal())
        text = self._extract(pdf)
        self.assertIn("VETO", text)
        self.assertIn("一票否决 ×1", text)
        # Phase 5 第四轮 A1：否决态主词不许再是三档词——"PASS · 否决态"必须零命中
        self.assertNotIn("PASS · 否决态", text)
        self.assertNotIn("结论为否决态", text)

    def test_uspap_and_fee_annotation(self):
        d = uw_commercial.example_39_main()
        r = uw_commercial.compute_all(d)
        # 构造费用未填场景：总费用 $0 必须标注"待接/未填"
        r["historical"]["total_expenses"] = 0
        r["proforma"]["total_expenses"] = 0
        pdf = pdf_uw.build_uw_pdf(r, "enhanced", deal=self._deal())
        text = self._extract(pdf)
        self.assertIn("非 USPAP 合规评估报告", text)
        self.assertIn("待接/未填", text)

    def test_no_meaningless_roi_row(self):
        # market cap 缺失 → projected_resale 算不出 → ROI/净收益必须"—"不许 -333%
        d = uw_commercial.example_39_main()
        d["analysis"]["market_cap_rate"] = 0
        r = uw_commercial.compute_all(d)
        pdf = pdf_uw.build_uw_pdf(r, "enhanced", deal=self._deal())
        text = self._extract(pdf)
        self.assertNotIn("-333", text)
        self.assertNotIn("333.33%", text)

    def test_assumptions_chinese_labels(self):
        d = uw_commercial.example_39_main()
        r = uw_commercial.compute_all(d)
        pdf = pdf_uw.build_uw_pdf(r, "enhanced", deal=self._deal())
        text = self._extract(pdf)
        self.assertIn("交易结构", text)
        self.assertIn("收购价/要价", text)
        self.assertNotIn("exitStrategy", text)

    def test_bank_dscr_in_pdf(self):
        d = uw_commercial.example_39_main()
        r = uw_commercial.compute_all(d)
        pdf = pdf_uw.build_uw_pdf(r, "enhanced", deal=self._deal())
        text = self._extract(pdf)
        self.assertIn("银行口径", text)
        self.assertIn("0.93", text)


if __name__ == "__main__":
    unittest.main()
