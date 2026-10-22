"""v1 金标准（golden master）公共库：用例重放 + 输出归一化。

两层锁定（v2 蓝图 §8.1：v1 与 v2 对同一批交易的所有输出数字一致，差异 = 0）：
- 函数层：纯计算模块的公开函数，输入取自 v1 全量测试套件实际传入的参数（录制所得）。
- API 层：HTTP 端点契约，输入 = 测试套件实际请求 + 交易矩阵（每笔样例交易 × 每个确定性端点）。

归一化只抹掉“生成时间”类字段与随时间变化的文本；数字一律原样（float 按 repr 精确比较）。
PDF 比对抽取出的文字，Excel 比对全部单元格值与公式，PNG 比对字节哈希。
外部网络一律断开（httpx 传输层抛 ConnectError），依赖外部数据的端点锁定其确定性降级输出。
"""

from __future__ import annotations

import base64
import contextlib
import gzip
import hashlib
import importlib
import io
import json
import math
import os
import re
import socket
from pathlib import Path
from typing import Any, Iterator

HERE = Path(__file__).resolve().parent
GOLDEN_DIR = HERE / "golden_v1"

# 时间戳类文本：2026-10-21 / 2026-10-21T19:37:00 / 2026/10/21 19:33:49 / 10-21 19:37 等
_TS_PATTERNS = [
    re.compile(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2}(\.\d+)?)?(Z|[+-]\d{2}:?\d{2})?"),
    re.compile(r"\d{4}/\d{1,2}/\d{1,2}( \d{1,2}:\d{2}(:\d{2})?)?"),
    re.compile(r"\d{4}-\d{2}-\d{2}"),
    re.compile(r"\d{4}年\d{1,2}月\d{1,2}日"),
    re.compile(r"\b\d{1,2}:\d{2}(:\d{2})?\b"),
    # 文件名里的紧凑日期 20261021（导出文件名带当天日期）
    re.compile(r"(?<!\d)20\d{2}(0[1-9]|1[0-2])(0[1-9]|[12]\d|3[01])(?!\d)"),
]
# 以 epoch 秒/毫秒存的时间字段（键名判断）
_TIME_KEYS = {"created_at", "updated_at", "generated_at", "fetched_at", "at", "ts", "timestamp",
              "as_of", "asof", "computed_at"}


def norm_text(s: str) -> str:
    for p in _TS_PATTERNS:
        s = p.sub("<TS>", s)
    return s


def norm(obj: Any, key: str | None = None) -> Any:
    """递归归一化 JSON 样对象。"""
    if key in _TIME_KEYS and isinstance(obj, (int, float, str)) and not isinstance(obj, bool):
        return "<TS>"
    if isinstance(obj, dict):
        return {k: norm(v, k) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [norm(v) for v in obj]
    if isinstance(obj, str):
        return norm_text(obj)
    if isinstance(obj, float):
        if math.isnan(obj):
            return "<NaN>"
        if math.isinf(obj):
            return "<Inf>" if obj > 0 else "<-Inf>"
        return obj
    if isinstance(obj, (bytes, bytearray)):
        return {"__bytes_sha256": hashlib.sha256(bytes(obj)).hexdigest(), "len": len(obj)}
    if isinstance(obj, io.BytesIO):
        return norm(obj.getvalue())
    if hasattr(obj, "model_dump"):
        return norm(obj.model_dump())
    return obj


def pdf_digest(data: bytes) -> dict:
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    return {"pages": len(r.pages),
            "text": [norm_text(p.extract_text() or "") for p in r.pages]}


def xlsx_digest(data: bytes) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    out = {}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    v = c.value
                    cells[c.coordinate] = norm_text(v) if isinstance(v, str) else (
                        norm_text(v.isoformat()) if hasattr(v, "isoformat") else v)
        out[ws.title] = cells
    return out


def response_digest(resp) -> dict:
    ctype = resp.headers.get("content-type", "")
    d: dict[str, Any] = {"status": resp.status_code}
    if "application/json" in ctype:
        d["json"] = norm(resp.json())
    elif "application/pdf" in ctype:
        d["pdf"] = pdf_digest(resp.content)
    elif "spreadsheetml" in ctype:
        d["xlsx"] = xlsx_digest(resp.content)
    elif ctype.startswith("text/"):
        d["text"] = norm_text(resp.text)
    else:
        d["bytes_sha256"] = hashlib.sha256(resp.content).hexdigest()
    disp = resp.headers.get("content-disposition")
    if disp:
        d["disposition"] = norm_text(disp)
    return d


# ---------------------------------------------------------------- 断网
@contextlib.contextmanager
def network_blocked() -> Iterator[None]:
    import httpx

    def _deny(*_a, **_k):
        raise httpx.ConnectError("golden: network blocked")

    def _deny_dns(*_a, **_k):
        raise socket.gaierror("golden: network blocked")

    saved = (httpx.HTTPTransport.handle_request, httpx.AsyncHTTPTransport.handle_async_request,
             socket.getaddrinfo)
    httpx.HTTPTransport.handle_request = _deny  # type: ignore[method-assign]

    async def _adeny(*_a, **_k):
        raise httpx.ConnectError("golden: network blocked")

    httpx.AsyncHTTPTransport.handle_async_request = _adeny  # type: ignore[method-assign]
    socket.getaddrinfo = _deny_dns
    env_saved = {k: os.environ.pop(k) for k in list(os.environ) if k.startswith("TOPHAP_")}
    try:
        yield
    finally:
        (httpx.HTTPTransport.handle_request, httpx.AsyncHTTPTransport.handle_async_request,
         socket.getaddrinfo) = saved
        os.environ.update(env_saved)


# ---------------------------------------------------------------- 读写
def read_jsonl_gz(path: Path) -> list[dict]:
    with gzip.open(path, "rt", encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def write_jsonl_gz(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    # mtime=0：同内容同字节，git 不产生伪差异
    with open(path, "wb") as raw, gzip.GzipFile(fileobj=raw, mode="wb", mtime=0) as gz:
        for r in rows:
            gz.write((json.dumps(r, ensure_ascii=False, sort_keys=True) + "\n").encode("utf-8"))


# ---------------------------------------------------------------- 重放
def run_fn_case(case: dict) -> Any:
    mod = importlib.import_module("app." + case["mod"])
    fn = getattr(mod, case["fn"])
    args = json.loads(json.dumps(case["args"]))  # 深拷贝：被测函数可能原地改参数
    try:
        out = fn(*args["a"], **args["kw"])
    except Exception as e:  # noqa: BLE001 — 异常也是 v1 行为的一部分，一并锁定
        return {"__raises": type(e).__name__, "msg": norm_text(str(e))}
    return norm(out)


def run_api_cases(cases: list[dict], client) -> list[dict]:
    """按顺序重放 API 用例。{pid:N} 占位引用第 N 个 create 用例返回的 id。"""
    created: list[int] = []
    out = []
    for c in cases:
        url = c["url"]
        for i, pid in enumerate(created):
            url = url.replace("{pid:%d}" % i, str(pid))
        kw: dict[str, Any] = {}
        if c.get("json") is not None:
            kw["json"] = c["json"]
        if c.get("files"):
            kw["files"] = {k: (v["name"], base64.b64decode(v["b64"]), v["type"])
                           for k, v in c["files"].items()}
        resp = client.request(c["method"], url, **kw)
        if c.get("capture_id"):
            created.append(resp.json()["id"])
        out.append(response_digest(resp))
    return out
