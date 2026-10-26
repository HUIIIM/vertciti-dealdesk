"""文件：智能录入（PDF/截图/文本）、租金表模板与解析、PDF/Excel/Tear Sheet 导出。"""

from __future__ import annotations

import io

from fastapi import APIRouter, File, HTTPException, Request, UploadFile
from fastapi.responses import Response

from app.api.common import (
    _check_upload_size,
    _parse_rr_date,
    _WbIntakeReq,
)
from app.domain import (
    validators,
)
from app.services import image_intake, pdf_intake, research_pipeline

router = APIRouter(tags=["documents"])


@router.post("/api/wb/intake/gate")
def wb_intake_gate(payload: dict):
    """Intake 数据质量门：对任意字段集合跑门。

    payload: {deal_type: commercial|residential, fields: {flat_key: value} 或
              [{key,value,...}], tenants: [...], text: 原始文本(可选，用于文档内矛盾检测)}
    fail → blocked（拦截进 underwriting）；warn → 标黄复核。
    """
    deal_type = str(payload.get("deal_type") or "commercial").lower()
    if deal_type not in ("commercial", "residential"):
        raise HTTPException(400, f"unknown deal_type: {deal_type!r}")
    fields = payload.get("fields") or {}
    tenants = payload.get("tenants") or []
    text = str(payload.get("text") or "")
    try:
        return validators.run_quality_gate(deal_type, fields, tenants=tenants, raw_text=text)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, f"质量门运行失败：{str(e)[:200]}")


@router.post("/api/wb/intake/run")
def wb_intake_run(req: _WbIntakeReq):
    """智能搜集：纯地址或房源链接 → 全网公开信息自动搜集.

    慢接口（多次公网抓取，约 30-90 秒）。每个字段带来源/抓取时间/可信度；
    抓不到标"需手动补"；卖方口径标"卖方口径、待验证"。
    """
    if req.mode not in ("address", "url"):
        raise HTTPException(400, f"unknown mode: {req.mode}")
    try:
        return research_pipeline.run(req.mode, req.text)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, f"搜集失败：{str(e)[:200]}")


@router.post("/api/wb/intake/pdf")
async def wb_intake_pdf(request: Request, file: UploadFile = File(...)):
    """PDF intake：上传房源 flyer/OM/卖方材料 → 文本提取 → 结构化字段.

    全部字段标"卖方材料口径、待独立验证"。
    """
    name = file.filename or "upload.pdf"
    if not name.lower().endswith(".pdf"):
        raise HTTPException(400, "只接受 PDF 文件")
    _check_upload_size(request, 20 * 1024 * 1024, "PDF")
    data = await file.read()
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(400, "PDF 超过 20MB 上限")
    if len(data) < 100:
        raise HTTPException(400, "文件过小或为空")
    if not data[:5] == b"%PDF-":
        raise HTTPException(400, "文件不是有效的 PDF（缺少 PDF 魔数），请检查文件")
    path = pdf_intake.save_upload(data, name)
    try:
        return pdf_intake.run_pdf_upload(path, name)
    finally:
        try:
            import os
            os.remove(path)
        except Exception:
            pass


@router.post("/api/wb/intake/pdf-ocr-text")
def wb_intake_pdf_ocr_text(payload: dict):
    """住宅扫描版 PDF OCR 文本 intake。

    扫描 PDF 无文本层时 /api/wb/intake/pdf 返回 needs_ocr；浏览器端
    pdf.js 渲染 + Tesseract.js OCR 后把文本发到这里，走同一抽取引擎，
    全部标"扫描版 PDF OCR 提取、卖方材料口径、待独立验证"。
    """
    text = str((payload or {}).get("text") or "")
    filename = str((payload or {}).get("filename") or "scanned.pdf")
    if not text.strip():
        raise HTTPException(400, "OCR 文本为空，请换更清晰的扫描件重试")
    if len(text) > 500_000:
        raise HTTPException(400, "OCR 文本过长（上限约 500KB）")
    try:
        d = pdf_intake.run_pdf_ocr_text(text, filename)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, f"提取失败：{str(e)[:200]}")
    if d.get("error"):
        raise HTTPException(422, d["error"])
    return d


# ---------------- 商业核保 PDF intake：OM/flyer → 商业字段 → 自动填表 ----------------

@router.post("/api/uw-commercial/intake/pdf")
async def uw_commercial_intake_pdf(request: Request, file: UploadFile = File(...)):
    """商业 PDF intake：上传 OM/flyer/卖方材料 → 提取商业核保字段.

    返回 fields[]（key 直接对应 uw-commercial.html 的 data-in 路径）
    与 tenants[]（租约表行），全部标"卖方材料口径、待独立验证"。
    """
    name = file.filename or "upload.pdf"
    if not name.lower().endswith(".pdf"):
        raise HTTPException(400, "只接受 PDF 文件")
    _check_upload_size(request, 20 * 1024 * 1024, "PDF")
    data = await file.read()
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(400, "PDF 超过 20MB 上限")
    if len(data) < 100:
        raise HTTPException(400, "文件过小或为空")
    if not data[:5] == b"%PDF-":
        raise HTTPException(400, "文件不是有效的 PDF（缺少 PDF 魔数），请检查文件")
    path = pdf_intake.save_upload(data, name)
    try:
        return pdf_intake.run_commercial_pdf_upload(path, name)
    finally:
        try:
            import os
            os.remove(path)
        except Exception:
            pass


@router.post("/api/uw-commercial/intake/pdf-ocr-text")
def uw_commercial_intake_pdf_ocr_text(payload: dict):
    """商业扫描版 PDF OCR 文本 intake。

    扫描 PDF 无文本层时 /api/uw-commercial/intake/pdf 返回 needs_ocr；
    浏览器端 pdf.js 渲染 + Tesseract.js OCR 后把文本发到这里，与普通
    PDF intake 共用同一抽取引擎；返回 fields[]/tenants[]/quality_gate，
    全部标"扫描版 PDF OCR 提取、卖方材料口径、待独立验证"。
    """
    text = str((payload or {}).get("text") or "")
    filename = str((payload or {}).get("filename") or "scanned.pdf")
    if not text.strip():
        raise HTTPException(400, "OCR 文本为空，请换更清晰的扫描件重试")
    if len(text) > 500_000:
        raise HTTPException(400, "OCR 文本过长（上限约 500KB）")
    try:
        d = pdf_intake.run_commercial_pdf_ocr_text(text, filename)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, f"提取失败：{str(e)[:200]}")
    if d.get("error"):
        raise HTTPException(422, d["error"])
    return d


# ---------------- 商业核保粘贴文本 intake：粘贴 OM/flyer 文本 → 商业字段 ----------------

@router.post("/api/uw-commercial/intake/text")
def uw_commercial_intake_text(payload: dict):
    """商业粘贴文本 intake：粘贴 OM/flyer/卖方材料文本
    → 提取商业核保字段（价格/面积/NOI/租户）。

    与 PDF intake 共用同一抽取引擎，返回形状一致（fields[] 直接对应
    uw-commercial.html 的 data-in 路径 + tenants[] + quality_gate），
    全部标"卖方材料口径、待独立验证"。
    """
    text = str((payload or {}).get("text") or "")
    if not text.strip():
        raise HTTPException(400, "文本为空，请粘贴 OM / flyer 文本后再提取")
    if len(text) > 500_000:
        raise HTTPException(400, "文本过长（上限约 500KB），请分段粘贴")
    try:
        return pdf_intake.run_commercial_text_upload(text)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, f"提取失败：{str(e)[:200]}")


# ---------------- 商业核保截图 intake：房源截图 → OCR → 商业字段 → 自动填表 ----------------

@router.post("/api/uw-commercial/intake/image-text")
def uw_commercial_intake_image_text(payload: dict):
    """商业截图 OCR 文本 intake。

    浏览器端 Tesseract.js 对截图做 OCR 后把文本发到这里，与 PDF intake
    共用同一抽取引擎；返回 fields[]/tenants[]/quality_gate，全部标
    "截图 OCR 提取、卖方材料口径、待独立验证"。
    """
    text = str((payload or {}).get("text") or "")
    filename = str((payload or {}).get("filename") or "screenshot.png")
    if not text.strip():
        raise HTTPException(400, "OCR 文本为空，请换更清晰的截图重试")
    if len(text) > 500_000:
        raise HTTPException(400, "OCR 文本过长（上限约 500KB）")
    try:
        d = pdf_intake.run_commercial_image_ocr_text(text, filename)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, f"提取失败：{str(e)[:200]}")
    if d.get("error"):
        raise HTTPException(422, d["error"])
    return d


@router.post("/api/uw-commercial/intake/image")
async def uw_commercial_intake_image(request: Request, file: UploadFile = File(...)):
    """商业截图直传 intake：服务端 OCR（tesseract 可用时）→ 商业字段。

    生产 serverless 无 tesseract 时返回 422 + error_code=ocr_unavailable_server，
    前端据此走浏览器端 OCR 路径，绝不假装排队提取（旧 pending 轮询端点
    /api/wb/intake/image/result/ 从未存在，商业页旧截图流程在生产不可用）。
    """
    name = file.filename or "screenshot.png"
    if not name.lower().endswith((".png", ".jpg", ".jpeg", ".webp")):
        raise HTTPException(400, "只接受截图图片（png/jpg/webp）")
    _check_upload_size(request, 15 * 1024 * 1024, "截图")
    data = await file.read()
    if len(data) > 15 * 1024 * 1024:
        raise HTTPException(400, "截图超过 15MB 上限")
    # 魔数校验与住宅截图 intake（image_intake.ALLOWED_EXT）共用同一规则
    magic_ok = (
        data[:8] == b"\x89PNG\r\n\x1a\n"
        or data[:2] == b"\xff\xd8"
        or (data[:4] == b"RIFF" and data[8:12] == b"WEBP")
    )
    if not magic_ok:
        raise HTTPException(400, "文件不是有效的图片")
    d = pdf_intake.run_commercial_image_upload(data, name)
    if d.get("error"):
        raise HTTPException(422, d["error"])
    return d


# ---------------- 截图 intake：房源页截图 → pending 队列 → cron 视觉提取 ----------------

@router.post("/api/wb/intake/image")
async def wb_intake_image(request: Request, file: UploadFile = File(...)):
    """截图 intake：上传房源页截图（png/jpg/webp）→ 存入 pending 队列.

    视觉提取由定时 watcher 完成（约10分钟内），结果经
    GET /api/wb/intake/pending 查询、GET /api/wb/intake/extracted/{task_id} 取回。
    全部字段标"截图提取、待验证"。
    """
    name = file.filename or "screenshot.png"
    _check_upload_size(request, 15 * 1024 * 1024, "截图")
    data = await file.read()
    try:
        return image_intake.save_screenshot(data, name)
    except ValueError as e:
        raise HTTPException(400, str(e))


@router.get("/api/wb/intake/pending")
def wb_intake_pending():
    """列出截图提取任务（pending/done/failed）。"""
    return {"tasks": image_intake.list_tasks()}


@router.get("/api/wb/intake/extracted/{task_id}")
def wb_intake_extracted(task_id: str):
    """取回某截图任务的视觉提取结果（renderIntake 兼容格式，附 quality_gate）。"""
    try:
        return image_intake.get_extracted(task_id)
    except FileNotFoundError:
        raise HTTPException(404, "提取结果尚未生成，请稍后刷新")
    except ValueError as e:
        raise HTTPException(400, str(e))


@router.post("/api/wb/intake/image-text")
def wb_intake_image_text(payload: dict):
    """住宅截图 OCR 文本 intake。

    浏览器端 OCR 后把文本发到这里，走住宅抽取引擎；返回
    fields[]/quality_gate（residential-v1），全部标"截图 OCR 提取、
    卖方材料口径、待独立验证"。与商业截图路同构，补齐住宅截图门覆盖。
    """
    text = str((payload or {}).get("text") or "")
    filename = str((payload or {}).get("filename") or "screenshot.png")
    if not text.strip():
        raise HTTPException(400, "OCR 文本为空，请换更清晰的截图重试")
    if len(text) > 500_000:
        raise HTTPException(400, "OCR 文本过长（上限约 500KB）")
    try:
        d = pdf_intake.run_residential_image_ocr_text(text, filename)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, f"提取失败：{str(e)[:200]}")
    if d.get("error"):
        raise HTTPException(422, d["error"])
    return d


@router.get("/api/uw/rentroll/template.xlsx")
def rentroll_template():
    """rent roll 导入模板（Excel）：逐租户行，用户填完上传."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "rent_roll"
    headers = ["单元", "租户", "面积SF", "月租金$", "租约开始(YYYY-MM-DD)",
               "租约结束(YYYY-MM-DD)", "年递增%", "费用结构(NNN/Gross/ModGross)"]
    ws.append(headers)
    ws.append(["101", "示例租户", 1200, 3500, "2024-01-01", "2027-12-31", 3, "NNN"])
    for i, w in enumerate([10, 20, 12, 12, 22, 22, 10, 26], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=dealdesk_rentroll_template.xlsx"},
    )


@router.post("/api/uw/rentroll/parse")
async def rentroll_parse(request: Request, file: UploadFile | None = File(None)):
    """rent roll 导入：上传 Excel 模板 或 body 粘贴 {"rows": [...]}.

    返回 uw 租户行格式（suite/tenant/sf/monthly_rent/lease_start/lease_end/
    escalation/expense_structure），可直接塞进 compute input.tenants。
    无 rent roll 时商业 verdict 最高 HOLD（由 verdict 引擎执行）。

    P0-13：模板示例行（"示例租户"）自动跳过并计数；日期解析失败只清空该字段
    （计入 bad_dates），不静默吞掉租户行。
    """
    rows: list[dict] = []
    skipped_example = 0  # P0-13：跳过的模板示例行数
    if file is not None:
        _check_upload_size(request, 5 * 1024 * 1024, "rent roll")
        data = await file.read()
        if len(data) > 5 * 1024 * 1024:
            raise HTTPException(400, "文件超过 5MB 上限")
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), data_only=True)
            ws = wb.active
            vals = list(ws.iter_rows(values_only=True))
        except Exception as e:  # noqa: BLE001
            raise HTTPException(400, f"Excel 解析失败：{str(e)[:120]}")
        if len(vals) < 2:
            raise HTTPException(400, "模板为空（至少需要表头＋1 行）")
        header = [str(h or "").strip() for h in vals[0]]
        def col(*names):
            for n in names:
                if n in header:
                    return header.index(n)
            return None
        c_suite = col("单元", "suite"); c_tenant = col("租户", "tenant")
        c_sf = col("面积SF", "sf"); c_rent = col("月租金$", "月租金", "monthly_rent")
        c_start = col("租约开始(YYYY-MM-DD)", "租约开始", "lease_start")
        c_end = col("租约结束(YYYY-MM-DD)", "租约结束", "lease_end")
        c_esc = col("年递增%", "escalation"); c_exp = col("费用结构(NNN/Gross/ModGross)", "费用结构")
        for line in vals[1:]:
            if not any(line):
                continue
            get = lambda c: line[c] if c is not None and c < len(line) else ""
            tenant_name = str(get(c_tenant) or "").strip()
            # P0-13：模板示例行不是真实租户，直接跳过
            if tenant_name == "示例租户" or tenant_name.startswith("示例"):
                skipped_example += 1
                continue
            rows.append({
                "suite": str(get(c_suite) or ""), "tenant": tenant_name,
                "sf": get(c_sf) or 0, "monthly_rent": get(c_rent) or 0,
                "lease_start": _parse_rr_date(get(c_start)),
                "lease_end": _parse_rr_date(get(c_end)),
                "escalation": str(get(c_esc) or ""),
                "expense_structure": str(get(c_exp) or ""),
            })
    else:
        body = await request.json()
        for r0 in (body.get("rows") or []):
            if isinstance(r0, dict):
                tn = str(r0.get("tenant", "")).strip()
                if tn == "示例租户" or tn.startswith("示例"):
                    skipped_example += 1
                    continue
                rows.append(r0)
    tenants = []
    bad_dates = 0  # P0-13：日期解析失败的租户行数（字段置空，不丢行）
    for r0 in rows:
        try:
            ls = _parse_rr_date(r0.get("lease_start", ""))
            le = _parse_rr_date(r0.get("lease_end", ""))
            if (r0.get("lease_start") or r0.get("lease_end")) and not (ls or le):
                bad_dates += 1
            tenants.append({
                "suite": str(r0.get("suite", "")),
                "tenant": str(r0.get("tenant", "")),
                "sf": float(r0.get("sf") or 0),
                "monthly_rent": float(r0.get("monthly_rent") or 0),
                "lease_start": ls,
                "lease_end": le,
                "escalation": str(r0.get("escalation", "")),
                "expense_structure": str(r0.get("expense_structure", "")),
            })
        except (TypeError, ValueError):
            continue
    return {
        "tenants": tenants,
        "count": len(tenants),
        "skipped_example": skipped_example,
        "bad_dates": bad_dates,
        "total_monthly": round(sum(t["monthly_rent"] for t in tenants), 2),
        "note": "租约明细待补充" if not tenants else "已导入，可进 compute_all().tenants",
    }
