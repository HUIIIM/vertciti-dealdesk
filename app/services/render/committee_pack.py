"""投委会包渲染：PDF（备忘录 + 关键图表 + 风险清单 + 投票记录）与 Excel，vertciti 品牌规范。

入参是 services.v2.committee.pack_data() 组装好的纯数据；本模块只管排版。
数字口径：千分位、负数括号 + 红色（色盲友好的双重标示）、未知一律“—”。
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.graphics.shapes import Drawing, Line, Rect, String
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core import brand
from app.services.render import pdf_report

JSON = dict[str, Any]
INK, STONE, LINE, MIST = (HexColor(brand.INK), HexColor(brand.STONE), HexColor(brand.LINE), HexColor(brand.MIST))
RED, GREEN, ORANGE, AMBER = (HexColor(brand.RED), HexColor(brand.GREEN), HexColor(brand.ORANGE),
                             HexColor(brand.AMBER_TEXT))


# ---------------------------------------------------------------- 数字格式
def fmt(v: Any, unit: str) -> str:
    if v is None:
        return "—"
    if unit in ("usd", "usd_month"):
        s = f"${abs(v):,.0f}"
        return f"({s})" if v < 0 else s
    if unit == "pct":
        return f"({abs(v):.2f}%)" if v < 0 else f"{v:.2f}%"
    return f"({abs(v):.2f}×)" if v < 0 else f"{v:.2f}×"


def _neg(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and v < 0


# ---------------------------------------------------------------- PDF
def _styles() -> dict[str, ParagraphStyle]:
    pdf_report._register_fonts()
    f = pdf_report.FONT
    return {
        "h1": ParagraphStyle("h1", fontName=f, fontSize=20, leading=26, textColor=INK, spaceAfter=4),
        "eyebrow": ParagraphStyle("eb", fontName=f, fontSize=8.5, leading=11, textColor=STONE),
        "h2": ParagraphStyle("h2", fontName=f, fontSize=12.5, leading=17, textColor=INK, spaceBefore=10, spaceAfter=4),
        "body": ParagraphStyle("b", fontName=f, fontSize=9.5, leading=14, textColor=INK),
        "small": ParagraphStyle("s", fontName=f, fontSize=8, leading=11, textColor=STONE),
        "cell": ParagraphStyle("c", fontName=f, fontSize=8.5, leading=11.5, textColor=INK),
        "num": ParagraphStyle("n", fontName=f, fontSize=8.5, leading=11.5, textColor=INK, alignment=2),
    }


def _esc(t: Any) -> str:
    return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _table(rows: list[list[Any]], widths: list[float], header: bool = True) -> Table:
    t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
    style = [("FONTNAME", (0, 0), (-1, -1), pdf_report.FONT), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
             ("LINEBELOW", (0, 0), (-1, -1), 0.4, LINE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]
    if header:
        style += [("LINEBELOW", (0, 0), (-1, 0), 1.2, INK), ("TEXTCOLOR", (0, 0), (-1, 0), STONE)]
    t.setStyle(TableStyle(style))
    return t


def _section(st: dict[str, ParagraphStyle], text: str) -> Table:
    t = Table([[Paragraph(_esc(text), st["h2"])]], colWidths=[174 * mm])
    t.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 3, INK), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    return t


def _tornado(rt: JSON, width: float = 174 * mm) -> Drawing:
    top = rt.get("top") or []
    h = 22 + 20 * max(len(top), 1)
    d = Drawing(width, h)
    if not top:
        d.add(String(0, h / 2, "红队面板：数据不足，暂无图", fontName=pdf_report.FONT, fontSize=9, fillColor=STONE))
        return d
    label_w, mid = 70 * mm, 70 * mm + (width - 70 * mm) / 2
    span = max(max(abs(r["minus"]["delta_pct_points"]), abs(r["plus"]["delta_pct_points"])) for r in top) or 1
    scale = ((width - label_w) / 2 - 6) / span
    d.add(Line(mid, 4, mid, h - 14, strokeColor=INK, strokeWidth=0.8))
    d.add(String(mid, h - 10, f"基准 {rt['metric_label']} {fmt(rt.get('base'), 'pct')}", fontName=pdf_report.FONT,
                 fontSize=7.5, fillColor=STONE, textAnchor="middle"))
    for i, r in enumerate(top):
        y = h - 32 - i * 20
        d.add(String(0, y + 3, f"{r['label']}（±{rt['shock_pct']:g}%）", fontName=pdf_report.FONT, fontSize=8,
                     fillColor=INK))
        for side in ("minus", "plus"):
            dv = r[side]["delta_pct_points"]
            w = abs(dv) * scale
            x = mid if dv >= 0 else mid - w
            d.add(Rect(x, y, w, 11, fillColor=RED if dv < 0 else GREEN, strokeColor=None))
        if r.get("material"):
            d.add(String(width, y + 3, "关键", fontName=pdf_report.FONT, fontSize=7.5, fillColor=ORANGE,
                         textAnchor="end"))
    return d


def _footer(pack: JSON):  # noqa: ANN202 — reportlab 回调
    def draw(canvas, doc) -> None:  # noqa: ANN001
        canvas.saveState()
        canvas.setFillColor(INK)
        canvas.rect(0, A4[1] - 10 * mm, A4[0], 10 * mm, stroke=0, fill=1)
        canvas.setFillColor(HexColor(brand.PAPER))
        canvas.setFont(pdf_report.FONT, 9)
        canvas.drawString(18 * mm, A4[1] - 6.6 * mm, "vertciti · DealDesk 投委会包")
        canvas.setFillColor(ORANGE)
        canvas.rect(18 * mm, A4[1] - 10 * mm, 30 * mm, 1.2 * mm, stroke=0, fill=1)
        canvas.setFillColor(STONE)
        canvas.setFont(pdf_report.FONT, 7)
        canvas.drawString(18 * mm, 9 * mm, f"运行编号 {pack['run_id']} · 引擎 {pack['engine_version']} · "
                                           f"数据摘要 {pack['data_sha256'][:16]}")
        canvas.drawRightString(A4[0] - 18 * mm, 9 * mm, f"第 {doc.page} 页")
        canvas.restoreState()
    return draw


def build_pdf(pack: JSON) -> bytes:
    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm,
                            bottomMargin=16 * mm, title=f"投委会包 · {pack['deal']['name']}", author="vertciti DealDesk")
    deal, m, cat = pack["deal"], pack["metrics"], pack["metric_catalog"]
    story: list[Any] = [
        Paragraph(_esc(f"投委会包 · {deal['stage_label']} · 第 {pack['version']['version_no']} 版"), st["eyebrow"]),
        Paragraph(_esc(deal["name"]), st["h1"]),
        Paragraph(_esc(f"{deal.get('address') or '—'} · {'商业' if deal['track'] == 'commercial' else '住宅'} · "
                       f"生成 {pack['generated_at']}"), st["small"]),
        Spacer(1, 6)]
    # 结论卡
    kill = pack["kill"]
    kill_txt = ("已触发：" + "；".join(r["text"] for r in kill["triggered"])) if kill["triggered"] else (
        ("未知：" + "；".join(r["text"] for r in kill["unknown"])) if kill["unknown"] else "未触发")
    mb = pack.get("max_bid") or {}
    card = [[Paragraph(f"<b>{_esc(cat[k]['label'])}</b>", st["small"]) for k in pack["card_keys"]],
            [Paragraph(f"<font color='{brand.RED if _neg(m.get(k)) else brand.INK}' size=13>"
                       f"{_esc(fmt(m.get(k), cat[k]['unit']))}</font>", st["cell"]) for k in pack["card_keys"]]]
    story += [_table(card, [174 * mm / len(pack["card_keys"])] * len(pack["card_keys"]), header=False), Spacer(1, 4),
              Paragraph(_esc(f"最高出价：{fmt(mb.get('max_bid'), 'usd')}"
                             + (f"（卡在：{mb['binding']['label']}）" if mb.get("binding") else "")
                             + f" · 要价 {fmt(mb.get('asking'), 'usd')}"), st["body"]),
              Paragraph(f"淘汰条件：<font color='{brand.RED if kill['triggered'] else brand.INK}'>{_esc(kill_txt)}</font>",
                        st["body"]),
              Paragraph(_esc(f"核对进度：关键数字已复核 {pack['verification']['key_reviewed']}/"
                             f"{pack['verification']['key_total']} · 未裁决矛盾 {pack['verification']['open_conflicts']} · "
                             f"过期证据 {pack['verification']['stale_evidence']}"), st["body"])]
    # 自上一版起改了什么
    story.append(_section(st, "自上一版起改了什么"))
    diff = pack.get("diff")
    if not diff:
        story.append(Paragraph("这是第一版承销，无可对比的上一版。", st["body"]))
    else:
        story.append(Paragraph(_esc(f"对比第 {diff['a']['version_no']} 版（{diff['a']['label']}）：改动 "
                                    f"{diff['changed_inputs']} 项假设，{diff['changed_metrics']} 项结果随之变化。"),
                               st["small"]))
        rows: list[list[Any]] = [["假设路径", "上一版", "本版"]]
        for r in diff["inputs"][:25]:
            rows.append([Paragraph(_esc(r["path"]), st["cell"]), Paragraph(_esc(r["before"]), st["num"]),
                         Paragraph(_esc(r["after"]), st["num"])])
        story.append(_table(rows, [94 * mm, 40 * mm, 40 * mm]))
    # 关键图表
    story += [_section(st, "红队面板：最可能出错的 5 个假设"), _tornado(pack["redteam"]),
              Paragraph(_esc(pack["redteam"].get("note", "")), st["small"])]
    sc = pack.get("scenarios")
    if sc:
        rows = [["情景", "概率", sc["metric_label"], "亏损"]]
        for r in sc["scenarios"]:
            rows.append([r["name"], f"{r['probability']:g}%", fmt(r["primary"], "pct"),
                         "—" if r["loss"] is None else ("是" if r["loss"] else "否")])
        rows.append(["期望值", "100%", fmt(sc["expected"], "pct"),
                     "—" if sc["loss_probability_pct"] is None else f"亏损概率 {sc['loss_probability_pct']:g}%"])
        story += [_section(st, "概率加权结果"), _table(rows, [50 * mm, 30 * mm, 50 * mm, 44 * mm])]
    # 反方意见：上会前先看“最可能怎么失败”
    dv = pack.get("devils") or {}
    story += [PageBreak(), _section(st, "反方意见（系统按固定规则生成，不经 AI）")]
    if dv.get("points"):
        story += [Paragraph(_esc(f"{'■ 高' if p['severity'] == 'high' else '□ 中'} · {p['title']}：{p['text']}"), st["body"])
                  for p in dv["points"]]
    else:
        story += [Paragraph(_esc("未发现可由数据支持的反方论点（数据不足的论点不写）"), st["body"])]
    # 风险清单
    story += [_section(st, "风险清单")]
    risks = pack["risks"] or ["未发现需要上会讨论的风险项"]
    story += [Paragraph(_esc(f"· {r}"), st["body"]) for r in risks]
    # 关键指标全表
    rows = [["指标", "数值"]] + [[cat[k]["label"], Paragraph(
        f"<font color='{brand.RED if _neg(v) else brand.INK}'>{_esc(fmt(v, cat[k]['unit']))}</font>", st["num"])]
        for k, v in m.items() if k in cat]
    story += [_section(st, "全部标准指标"), _table(rows, [110 * mm, 64 * mm])]
    # 投票记录
    cm = pack["committee"]
    story += [_section(st, "投票记录"),
              Paragraph(_esc(f"结果：{ {'approved': '通过', 'rejected': '不通过', 'pending': '待定'}[cm['result']] }"
                             f"（{cm['why']}）· 赞成 {cm['approve']} · 反对 {cm['reject']} · 弃权 {cm['abstain']}"),
                        st["body"])]
    vrows: list[list[Any]] = [["投票人", "角色", "投票", "意见", "时间"]]
    for v in pack["votes"]:
        vrows.append([v["voter"], v.get("role_label", ""), {"approve": "赞成", "reject": "反对", "abstain": "弃权"}[v["vote"]],
                      Paragraph(_esc(v.get("comment") or ""), st["cell"]), v["created_at"][:16].replace("T", " ")])
    story.append(_table(vrows, [28 * mm, 18 * mm, 16 * mm, 80 * mm, 32 * mm]))
    if pack.get("decisions"):
        story.append(_section(st, "决策日志"))
        for dcs in pack["decisions"]:
            story.append(Paragraph(_esc(f"{dcs['created_at'][:10]} {dcs['verdict']} · {dcs['created_by']}：{dcs['rationale']}"),
                                   st["body"]))
    story += [Spacer(1, 10), Paragraph(_esc(
        "校验：本包内容摘要与运行编号已登记；把 PDF 上传到 /api/v2/committee/verify 即可核验是否原样。"
        "筛选辅助工具，不构成投资建议。"), st["small"])]
    doc.build(story, onFirstPage=_footer(pack), onLaterPages=_footer(pack))
    return buf.getvalue()


# ---------------------------------------------------------------- Excel
def build_xlsx(pack: JSON) -> bytes:
    wb = Workbook()
    hdr_fill, hdr_font = PatternFill("solid", fgColor=brand.INK[1:]), Font(bold=True, color=brand.PAPER[1:])
    red = Font(color=brand.RED[1:])

    def sheet(title: str, header: list[str], rows: list[list[Any]], first: bool = False) -> None:
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = hdr_fill, hdr_font
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "B2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, min(48, max(len(str(c.value or "")) for c in col) + 2))
        for rr in ws.iter_rows(min_row=2):
            for c in rr:
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    c.alignment = Alignment(horizontal="right")
                    c.number_format = "#,##0.00;[Red](#,##0.00)"
                    if c.value < 0:
                        c.font = red

    cat, m = pack["metric_catalog"], pack["metrics"]
    sheet("摘要", ["项目", "值"], [["交易", pack["deal"]["name"]], ["地址", pack["deal"].get("address") or ""],
                               ["阶段", pack["deal"]["stage_label"]], ["版本", pack["version"]["version_no"]],
                               ["运行编号", pack["run_id"]], ["引擎版本", pack["engine_version"]],
                               ["数据摘要 SHA-256", pack["data_sha256"]], ["生成时间", pack["generated_at"]],
                               ["最高出价", (pack.get("max_bid") or {}).get("max_bid")]], first=True)
    sheet("指标", ["指标", "单位", "数值"], [[cat[k]["label"], cat[k]["unit"], v] for k, v in m.items() if k in cat])
    sheet("红队面板", ["假设", "-10% 后", "+10% 后", "波动(bp)", "关键"],
          [[r["label"], r["minus"]["value"], r["plus"]["value"], r["swing_bps"], "是" if r["material"] else ""]
           for r in pack["redteam"].get("all", [])])
    if pack.get("scenarios"):
        sheet("情景", ["情景", "概率%", pack["scenarios"]["metric_label"], "亏损"],
              [[r["name"], r["probability"], r["primary"], r["loss"]] for r in pack["scenarios"]["scenarios"]])
    sheet("版本差异", ["路径", "类型", "上一版", "本版"],
          [[r["path"], r["kind"], str(r["before"]), str(r["after"])] for r in (pack.get("diff") or {}).get("inputs", [])])
    sheet("反方意见", ["级别", "论点", "说明"],
          [["高" if p["severity"] == "high" else "中", p["title"], p["text"]] for p in (pack.get("devils") or {}).get("points", [])])
    sheet("风险清单", ["风险"], [[r] for r in pack["risks"]])
    sheet("投票", ["投票人", "投票", "意见", "时间"],
          [[v["voter"], v["vote"], v.get("comment") or "", v["created_at"]] for v in pack["votes"]])
    sheet("数字护照", ["字段", "值", "来源", "页", "方式", "可信度", "复核", "新鲜度"],
          [[p["field"], str(p["value"]), p["source"], p["page"], p["method_label"], p["confidence_label"],
            p["review_label"], p["freshness"]["label"]] for p in pack["passports"]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
